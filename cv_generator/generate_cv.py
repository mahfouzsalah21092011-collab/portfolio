#!/usr/bin/env python3
"""
Generate Mahfouz Salah's RDD / MEP CV as an editable DOCX and PDF.

Requirements:
    pip install python-docx reportlab

Usage:
    python generate_cv.py
Outputs:
    Mahfouz_Salah_CV_RDD_MEP.docx
    Mahfouz_Salah_CV_RDD_MEP.pdf

Also synchronizes with project downloads directory:
    ../downloads/Mahfouz_Salah_CV.docx
    ../downloads/Mahfouz_Salah_CV.pdf
    ../downloads/Mahfouz_Salah_CV_RDD_MEP.docx
    ../downloads/Mahfouz_Salah_CV_RDD_MEP.pdf
"""
from pathlib import Path
import shutil
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import parse_xml
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable, KeepTogether

OUT = Path(__file__).resolve().parent
DOCX_PATH = OUT / "Mahfouz_Salah_CV_RDD_MEP.docx"
PDF_PATH = OUT / "Mahfouz_Salah_CV_RDD_MEP.pdf"

NAME = "MAHFOUZ SALAH MAHFOUZ SOLIMAN"
HEADLINE = "FACILITIES MANAGER | MEP SYSTEMS | RETAIL DESIGN & DELIVERY (RDD) | SHOPPING MALL OPERATIONS & TECHNICAL OVERSIGHT"
CONTACT_LOCATION_PHONE = "Cairo, Egypt   |   +2 01027629474   |   +2 01221833667"
EMAIL_ADDR = "Mahfouzsoliman2021@gmail.com"
EMAIL_URL = f"mailto:{EMAIL_ADDR}"
PORTFOLIO_URL = "https://portfolio.mahfouz-salah.workers.dev/"
PORTFOLIO_DISPLAY = "portfolio.mahfouz-salah.workers.dev"
LINKEDIN_URL = "https://www.linkedin.com/in/mahfouz-salah-suleiman-9058b2207/"
LINKEDIN_DISPLAY = "linkedin.com/in/mahfouz-salah-suleiman-9058b2207"

def add_docx_hyperlink(paragraph, url, text, color="146677", underline=True, size_pt=8, bold=False):
    part = paragraph.part
    r_id = part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hl = parse_xml(
        f'<w:hyperlink xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        f'r:id="{r_id}" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"/>'
    )
    r = parse_xml('<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>')
    rPr = parse_xml('<w:rPr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>')
    if color:
        c = parse_xml(f'<w:color xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:val="{color}"/>')
        rPr.append(c)
    if underline:
        u = parse_xml('<w:u xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:val="single"/>')
        rPr.append(u)
    if bold:
        b = parse_xml('<w:b xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>')
        rPr.append(b)
    if size_pt:
        sz = parse_xml(f'<w:sz xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:val="{int(size_pt * 2)}"/>')
        rPr.append(sz)
    r.append(rPr)
    t = parse_xml(f'<w:t xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">{text}</w:t>')
    r.append(t)
    hl.append(r)
    paragraph._p.append(hl)
    return hl

SUMMARY = (
    "Facilities and MEP professional with 15+ years of experience across electrical maintenance, "
    "facility operations, commercial mall development, and retail tenant fit-outs. Experienced in "
    "reviewing MEP designs, specifications and shop drawings; following up execution; coordinating "
    "with tenants, contractors and consultants; and supporting inspections, testing, snagging and "
    "handover. Targeting Retail Design & Delivery (RDD) / MEP leadership opportunities."
)

SKILLS = [
    "Retail Design & Delivery (RDD) support and tenant fit-out coordination",
    "MEP design and shop-drawing review; technical comments and compliance follow-up",
    "Consultant-side site oversight, execution monitoring, inspections and snagging",
    "Testing, commissioning, documentation and handover coordination",
    "Electrical systems, HVAC, firefighting, fire alarm, plumbing, drainage and pumps",
    "Facilities operations, preventive/corrective maintenance and service contracts",
    "Budgeting, utilities, contractor/tenant/consultant coordination",
    "Civil Defense / fire-life-safety coordination, HSE and work permits",
    "AutoCAD, Microsoft Excel, Word and PowerPoint",
]

EXPERIENCE = [
    {
        "company": "MEGA PLAN",
        "dates": "Oct 2024 – Present",
        "role": "Facilities / MEP Department Manager – Consultant-side Technical Oversight",
        "bullets": [
            "Provide consultant-side technical oversight for MEP execution and retail tenant fit-out works.",
            "Review tenant MEP drawings, shop drawings, design details and technical submissions; coordinate comments and follow-up.",
            "Monitor site execution, inspections, testing, snagging and handover readiness, coordinating with tenants, contractors and relevant stakeholders.",
            "Coordinate Civil Defense and fire/life-safety requirements as applicable to project works.",
            "60 Feddan Club Mall: Manage MEP technical reviews, base-building interfaces, tenant submittal audits (power, HVAC, plumbing) to protect central plant capacity, and Civil Defense fire/life-safety inspections.",
            "Support tenant fit-outs across East Plaza Mall, Kayan Plaza Mall, and Souq El Madina Mall (approximately 11 tenant fit-outs collectively).",
            "Review drawings for 20+ tenants collectively across Mall Arena (Sarai Compound), Mall Taj (Taj City Compound) and Mall Kinda, and follow up MEP execution.",
            "Review drawings for 10 tenants and follow up MEP execution at Sky Walk, 6th of October City.",
            "Supervised works involving 35+ tenants at Muse Mall, near Al Naib Al Aam Bridge.",
        ],
    },
    {
        "company": "AL-AMAL FOR DEVELOPMENT AND TRADE",
        "dates": "Jan 2017 – Sep 2024",
        "role": "Facilities Manager / Head of Facilities Department",
        "bullets": [
            "Managed facilities operations and maintenance for Cross 10 Mall (approximately 22,000 m²).",
            "Led a facilities team of approximately 25 staff and coordinated day-to-day service delivery.",
            "Coordinated tenant fit-outs, maintenance activities, shop drawings, material approvals, testing, snagging and handover.",
            "Managed maintenance budgets, service contracts, utilities, suppliers and contractor performance.",
            "Oversaw a maintenance contract valued at approximately EGP 150,000 and supported inventory/process improvements.",
            "Coordinated works and tenant requirements across mall operations while maintaining safety and service continuity.",
        ],
    },
    {
        "company": "HYPER ONE TRADING COMPANY",
        "dates": "Nov 2016 – Jan 2017",
        "role": "Facilities Management Supervisor",
        "bullets": [
            "Supervised day-to-day facilities activities and coordinated maintenance support and service requests.",
        ],
    },
    {
        "company": "ENGINEERING CONTRACTING OFFICE – RAMADAN & HUSSEIN ABU HAMID",
        "dates": "Nov 2014 – Aug 2016",
        "role": "Electrical / Equipment Maintenance Engineer",
        "bullets": [
            "Supported electrical and equipment maintenance activities, troubleshooting and operational reliability.",
        ],
    },
    {
        "company": "INTERNATIONAL CONTRACTING COMPANY (MUNTASER)",
        "dates": "Nov 2010 – Oct 2014",
        "role": "Electrical Maintenance Engineer",
        "bullets": [
            "Performed electrical maintenance and troubleshooting and supported safe, reliable equipment operation.",
        ],
    },
]

EDUCATION = [
    ("Bachelor of Industrial Education", "Electronics Technology Department, Helwan University | 2009"),
    ("Graduation Project", "Miniature laboratory system for voltage/current measurement | Excellent"),
]

DEVELOPMENT = [
    "Applied Supervisory Professional Diploma, AUC School of Continuing Education | Nov 2022 – Nov 2023 | 210 hours",
    "Professional development courses via LinkedIn Learning | 2023–2025",
]

LANGUAGES = "Arabic: Native  |  English: Good"


def build_docx(docx_out: Path = DOCX_PATH):
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.5)
    sec.bottom_margin = Inches(0.5)
    sec.left_margin = Inches(0.65)
    sec.right_margin = Inches(0.65)
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9)
    normal.font.color.rgb = RGBColor(45, 55, 65)
    normal.paragraph_format.space_after = Pt(1.5)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(NAME)
    r.bold = True
    r.font.size = Pt(18)
    r.font.color.rgb = RGBColor(20, 53, 76)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(HEADLINE)
    r.bold = True
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor(20, 102, 119)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(1)
    p.add_run(CONTACT_LOCATION_PHONE + "   |   ").font.size = Pt(8)
    add_docx_hyperlink(p, EMAIL_URL, EMAIL_ADDR, color="146677", underline=True, size_pt=8)

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_after = Pt(2.5)

    r_li_lbl = p2.add_run("LinkedIn: ")
    r_li_lbl.bold = True
    r_li_lbl.font.size = Pt(8)
    r_li_lbl.font.color.rgb = RGBColor(20, 53, 76)
    add_docx_hyperlink(p2, LINKEDIN_URL, LINKEDIN_DISPLAY, color="146677", underline=True, size_pt=8)

    r_sep = p2.add_run("   |   ")
    r_sep.font.size = Pt(8)
    r_sep.font.color.rgb = RGBColor(120, 135, 150)

    r_pf_lbl = p2.add_run("Portfolio: ")
    r_pf_lbl.bold = True
    r_pf_lbl.font.size = Pt(8)
    r_pf_lbl.font.color.rgb = RGBColor(20, 53, 76)
    add_docx_hyperlink(p2, PORTFOLIO_URL, PORTFOLIO_DISPLAY, color="146677", underline=True, size_pt=8)

    def heading(label):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(label.upper())
        r.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(20, 102, 119)
        return p

    heading("Professional Summary")
    doc.add_paragraph(SUMMARY)

    heading("Core Competencies")
    for item in SKILLS:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.18)
        p.paragraph_format.first_line_indent = Inches(-0.12)
        p.paragraph_format.space_after = Pt(1)
        p.add_run(item)

    heading("Professional Experience")
    for job in EXPERIENCE:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4.5)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(job["company"])
        r.bold = True
        r.font.color.rgb = RGBColor(20, 53, 76)
        r = p.add_run("  |  " + job["dates"])
        r.italic = True
        r.font.size = Pt(8)

        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(1.5)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(job["role"])
        r.bold = True

        for idx, bullet in enumerate(job["bullets"]):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.left_indent = Inches(0.18)
            p.paragraph_format.first_line_indent = Inches(-0.12)
            p.paragraph_format.space_after = Pt(1)
            if len(job["bullets"]) <= 2:
                p.paragraph_format.keep_with_next = True
            p.add_run(bullet)

    heading("Education")
    for title, detail in EDUCATION:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(title + " — ")
        r.bold = True
        p.add_run(detail)

    heading("Professional Development")
    for item in DEVELOPMENT:
        p = doc.add_paragraph(item, style="List Bullet")
        p.paragraph_format.space_after = Pt(1)

    heading("Tools & Languages")
    doc.add_paragraph("Tools: AutoCAD, Microsoft Excel, Word and PowerPoint.")
    doc.add_paragraph(LANGUAGES)

    docx_out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(docx_out))


def build_pdf(pdf_out: Path = PDF_PATH):
    pdf_out.parent.mkdir(parents=True, exist_ok=True)
    doc = SimpleDocTemplate(
        str(pdf_out),
        pagesize=A4,
        rightMargin=14 * mm,
        leftMargin=14 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
        title="Mahfouz Salah – Facilities & MEP CV",
        author="Mahfouz Salah Mahfouz Soliman",
    )
    styles = getSampleStyleSheet()
    styles.add(
        ParagraphStyle(
            name="CVName",
            parent=styles["Title"],
            fontName="Helvetica-Bold",
            fontSize=18,
            leading=20,
            alignment=TA_CENTER,
            textColor=colors.HexColor("#14354C"),
            spaceAfter=3,
        )
    )
    styles.add(
        ParagraphStyle(
            name="CVHeadline",
            parent=styles["Normal"],
            fontName="Helvetica-Bold",
            fontSize=8.5,
            leading=10.5,
            alignment=TA_CENTER,
            textColor=colors.HexColor("#146677"),
            spaceAfter=4,
        )
    )
    styles.add(
        ParagraphStyle(
            name="CVContact",
            parent=styles["Normal"],
            fontSize=7.5,
            leading=9.5,
            alignment=TA_CENTER,
            textColor=colors.HexColor("#374151"),
            spaceAfter=2,
        )
    )
    styles.add(
        ParagraphStyle(
            name="SectionCV",
            parent=styles["Heading2"],
            fontName="Helvetica-Bold",
            fontSize=9.5,
            leading=11.5,
            textColor=colors.HexColor("#146677"),
            spaceBefore=7.5,
            spaceAfter=2.5,
            keepWithNext=True,
        )
    )
    styles.add(
        ParagraphStyle(
            name="BodyCV",
            parent=styles["BodyText"],
            fontSize=8.5,
            leading=10.8,
            spaceAfter=2,
            textColor=colors.HexColor("#2D3741"),
        )
    )
    styles.add(
        ParagraphStyle(
            name="BulletCV",
            parent=styles["BodyText"],
            fontSize=8.2,
            leading=10.2,
            leftIndent=11,
            firstLineIndent=-7,
            bulletIndent=0,
            spaceAfter=2.0,
            textColor=colors.HexColor("#2D3741"),
        )
    )
    styles.add(
        ParagraphStyle(
            name="CompanyCV",
            parent=styles["BodyText"],
            fontName="Helvetica-Bold",
            fontSize=8.8,
            leading=10.5,
            textColor=colors.HexColor("#14354C"),
            spaceBefore=4.5,
            spaceAfter=1,
            keepWithNext=True,
        )
    )
    styles.add(
        ParagraphStyle(
            name="RoleCV",
            parent=styles["BodyText"],
            fontName="Helvetica-Bold",
            fontSize=8.0,
            leading=9.8,
            spaceAfter=2,
            keepWithNext=True,
        )
    )

    contact_p1 = (
        f"{CONTACT_LOCATION_PHONE} &nbsp;|&nbsp; "
        f"<a href='{EMAIL_URL}'><font color='#146677'><u>{EMAIL_ADDR}</u></font></a>"
    )
    contact_p2 = (
        f"<b>LinkedIn:</b> <a href='{LINKEDIN_URL}'><font color='#146677'><u>{LINKEDIN_DISPLAY}</u></font></a>"
        " &nbsp;|&nbsp; "
        f"<b>Portfolio:</b> <a href='{PORTFOLIO_URL}'><font color='#146677'><u>{PORTFOLIO_DISPLAY}</u></font></a>"
    )

    story = [
        Paragraph(NAME, styles["CVName"]),
        Paragraph(HEADLINE, styles["CVHeadline"]),
        Paragraph(contact_p1, styles["CVContact"]),
        Paragraph(contact_p2, styles["CVContact"]),
        Spacer(1, 2.5),
    ]

    def sec_flow(title):
        return [
            Paragraph(title.upper(), styles["SectionCV"]),
            HRFlowable(
                width="100%",
                thickness=0.5,
                color=colors.HexColor("#B9CDD2"),
                spaceAfter=3,
            ),
        ]

    story.extend(sec_flow("Professional Summary"))
    story.append(Paragraph(SUMMARY, styles["BodyCV"]))

    story.extend(sec_flow("Core Competencies"))
    for item in SKILLS:
        story.append(Paragraph("• " + item, styles["BulletCV"]))

    story.extend(sec_flow("Professional Experience"))
    for job in EXPERIENCE:
        job_flowables = [
            Paragraph(
                job["company"]
                + " &nbsp; <font name='Helvetica' size='7'>| "
                + job["dates"]
                + "</font>",
                styles["CompanyCV"],
            ),
            Paragraph(job["role"], styles["RoleCV"]),
        ]
        if len(job["bullets"]) <= 2:
            for b in job["bullets"]:
                job_flowables.append(Paragraph("• " + b, styles["BulletCV"]))
            story.append(KeepTogether(job_flowables))
        else:
            story.append(KeepTogether(job_flowables[:2]))
            for b in job["bullets"]:
                story.append(Paragraph("• " + b, styles["BulletCV"]))

    edu_block = sec_flow("Education")
    for title, detail in EDUCATION:
        edu_block.append(
            Paragraph("<b>" + title + ":</b> " + detail, styles["BodyCV"])
        )
    story.append(KeepTogether(edu_block))

    dev_block = sec_flow("Professional Development")
    for item in DEVELOPMENT:
        dev_block.append(Paragraph("• " + item, styles["BulletCV"]))
    story.append(KeepTogether(dev_block))

    tools_block = sec_flow("Tools & Languages")
    tools_block.append(
        Paragraph("Tools: AutoCAD, Microsoft Excel, Word and PowerPoint.", styles["BodyCV"])
    )
    tools_block.append(Paragraph(LANGUAGES, styles["BodyCV"]))
    story.append(KeepTogether(tools_block))

    doc.build(story)


def sync_downloads():
    downloads_dir = OUT.parent / "downloads"
    if downloads_dir.exists():
        # Update both default CV and RDD MEP CV in downloads
        shutil.copy2(DOCX_PATH, downloads_dir / "Mahfouz_Salah_CV.docx")
        shutil.copy2(PDF_PATH, downloads_dir / "Mahfouz_Salah_CV.pdf")
        shutil.copy2(DOCX_PATH, downloads_dir / "Mahfouz_Salah_CV_RDD_MEP.docx")
        shutil.copy2(PDF_PATH, downloads_dir / "Mahfouz_Salah_CV_RDD_MEP.pdf")
        print(f"Synchronized outputs to: {downloads_dir}")


if __name__ == "__main__":
    build_docx()
    build_pdf()
    print(f"Created: {DOCX_PATH}")
    print(f"Created: {PDF_PATH}")
    sync_downloads()
